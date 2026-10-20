"""Smoke tests that run on Linux without Chrome, Swift, or network access."""
import contextlib
import importlib
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'scripts'
sys.path.insert(0, str(SCRIPTS))

import audit_pipeline  # noqa: E402
import verify_common  # noqa: E402


class TestUrlNormalization(unittest.TestCase):
    def test_video_url(self):
        kind, key, canonical = verify_common.norm_url(
            'https://www.douyin.com/video/1234567890123456789'
        )
        self.assertEqual(kind, 'video')
        self.assertEqual(key, 'v:1234567890123456789')
        self.assertTrue(canonical.endswith('/video/1234567890123456789'))

    def test_share_video_url(self):
        kind, key, _ = verify_common.norm_url(
            'https://www.iesdouyin.com/share/video/1234567890123456789/?region=CN'
        )
        self.assertEqual(kind, 'video')
        self.assertEqual(key, 'v:1234567890123456789')

    def test_modal_id(self):
        kind, key, _ = verify_common.norm_url(
            'https://www.douyin.com/user/x?modal_id=1234567890123456789'
        )
        self.assertEqual(kind, 'video')
        self.assertEqual(key, 'v:1234567890123456789')

    def test_short_url(self):
        kind, key, _ = verify_common.norm_url('https://v.douyin.com/AbCdEf1/')
        self.assertEqual(kind, 'short')
        self.assertTrue(key.startswith('u:'))

    def test_unsupported(self):
        self.assertEqual(verify_common.norm_url('https://weixin.qq.com/sph/abc')[0], 'unsupported')
        self.assertEqual(verify_common.norm_url('not a url')[0], 'unsupported')


class TestOcrParsing(unittest.TestCase):
    def test_standard_labels(self):
        txt = ('播放量|1.5万|点赞量|100|评论量|20|收藏量|5|分享量|10|'
               '完播率|6.64%|作品状态|正常')
        finish, like, cmt, rep, fav, play, bad = audit_pipeline.parse_ocr(txt)
        self.assertEqual(finish, 6.64)
        self.assertEqual(like, 100)
        self.assertEqual(cmt, 20)
        self.assertEqual(rep, 10)
        self.assertEqual(fav, 5)
        self.assertEqual(play, 15000)
        self.assertFalse(bad)

    def test_vision_variant(self):
        finish, *_ = audit_pipeline.parse_ocr('元播率|5.07%')
        self.assertEqual(finish, 5.07)

    def test_excludes_5s(self):
        finish, *_ = audit_pipeline.parse_ocr('5s完播率|10%|完播率|8%')
        self.assertEqual(finish, 8)

    def test_deleted_signal(self):
        *_, bad = audit_pipeline.parse_ocr('作品状态|已删除')
        self.assertTrue(bad)


class TestRulesConfig(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()
        importlib.reload(audit_pipeline)

    def test_apply_rules(self):
        cfg = {
            'tiers': {
                'base100': {'label': '基础激励档'},
                'vip': {'play': 100000, 'interact': 0.01, 'finish': 10, 'label': 'VIP档'},
            },
            'incentive_tiers': {'200': 'base100', '900': 'vip'},
            'fee_tiers': {'800': 'vip'},
        }
        path = os.path.join(self.tmp.name, 'rules.json')
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False)
        audit_pipeline.apply_rules(path)
        self.assertIn('vip', audit_pipeline.TIER)
        self.assertEqual(audit_pipeline.INC_TIER[900], 'vip')
        self.assertEqual(audit_pipeline.FEE_TIER[800], 'vip')
        a3, a4 = audit_pipeline.rules_summary()
        self.assertIn('VIP档', a3)
        self.assertIn('VIP档', a4)

    def test_missing_field_raises(self):
        cfg = {'tiers': {'broken': {'play': 1}}, 'incentive_tiers': {}, 'fee_tiers': {}}
        path = os.path.join(self.tmp.name, 'bad.json')
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(cfg, f)
        with self.assertRaises(ValueError):
            audit_pipeline.apply_rules(path)

    def test_unknown_tier_reference_raises(self):
        cfg = {'tiers': {'base100': {'label': '基础'}},
               'incentive_tiers': {'100': 'missing'}, 'fee_tiers': {}}
        path = os.path.join(self.tmp.name, 'bad_ref.json')
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(cfg, f)
        with self.assertRaises(ValueError):
            audit_pipeline.apply_rules(path)


class TestTemplate(unittest.TestCase):
    def test_generate_template(self):
        from openpyxl import load_workbook
        import make_template
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, 'template.xlsx')
            old = sys.argv
            try:
                sys.argv = ['make_template.py', '--out', out, '--rows', '3']
                with contextlib.redirect_stdout(io.StringIO()):
                    make_template.main()
            finally:
                sys.argv = old
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ['验收链接', '审核规则', '填写说明'])
            ws = wb['验收链接']
            self.assertEqual(ws.max_column, 37)
            self.assertEqual(ws.cell(1, 1).value, '序号')
            self.assertEqual(ws.cell(1, 37).value, '人工确认项')


if __name__ == '__main__':
    unittest.main()
