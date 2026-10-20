"""贴门槛值钉死工具: 自动定位标签→裁剪放大(5x)→OCR复核
用法: swiftc -O bbox.swift -o /tmp/bboxbin   (先编译)
      python3 pin_value.py <源表.xlsx> <标签(如 完播率)> <序号...>
说明: 对每个序号取截图列(24/25/26)的URL/本地路径, 定位标签token坐标后裁剪其正下方区域并5x放大重OCR,
     用于「判定贴门槛」或「全页OCR未读到数值」的行做像素级裁决(历史批次经验)。
"""
import argparse
import json
import os
import subprocess
import sys
import openpyxl
from PIL import Image

ap = argparse.ArgumentParser(description='Pin a threshold value by cropping and re-OCRing it.')
ap.add_argument('xlsx')
ap.add_argument('label', help='标签文本，如 完播率')
ap.add_argument('seqs', nargs='+', type=int, help='序号列表')
args = ap.parse_args()
XLSX, LABEL, seqs = args.xlsx, args.label, args.seqs
BBOX='/tmp/bboxbin'; OCR=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ocr.swift')
CACHE=os.path.expanduser('~/.cache/yanshou_ocr.json')
if not os.path.exists(BBOX):
    print('缺少 bbox 工具, 正在编译...')
    subprocess.run(['swiftc','-O',os.path.join(os.path.dirname(os.path.abspath(__file__)), 'bbox.swift'),'-o',BBOX])
cache=json.load(open(CACHE)) if os.path.exists(CACHE) else {}
wb=openpyxl.load_workbook(XLSX, data_only=True); ws=wb['验收链接']
hdr=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
def find_label(img, label):
    p=subprocess.run([BBOX,img],capture_output=True,text=True)
    out=[]
    for line in p.stdout.splitlines():
        if '\t' not in line: continue
        xy,txt=line.split('\t',1)
        if label in txt and '5s' not in txt and '5S' not in txt:
            x,y,w,h=map(float,xy.split()); out.append((x,y,w,h,txt))
    return out
for seq in seqs:
    r=seq+1; v=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
    print(f'== 序号{seq}')
    for i in (23,24,25):
        s=str(v[i] or '').strip()
        if not s or s=='None': continue
        if os.path.exists(s): img=s
        else:
            ext=os.path.splitext(s.split('?')[0])[1] or '.jpg'
            img=f'/tmp/ys_pin/seq{seq}_s{i-22}{ext}'
            os.makedirs('/tmp/ys_pin', exist_ok=True)
            if not os.path.exists(img) or os.path.getsize(img)<1000:
                subprocess.run(['curl','-s','-o',img,'--max-time','40','-A',
                    'Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X)','-L',s])
        if not os.path.exists(img) or os.path.getsize(img)<1000:
            print(f'   截图{i-22}: 获取失败'); continue
        for x,y,w,h,txt in find_label(img, LABEL):
            im=Image.open(img); W,H=im.size
            x0=int(max(0,(x-0.04))*W); x1=int(min(1,(x+0.30))*W)
            top=int((1-(y+0.01))*H); bot=int((1-(y-0.055))*H)
            crop=im.crop((x0,top,x1,bot)); crop=crop.resize((crop.width*5,crop.height*5),Image.LANCZOS)
            cp='/tmp/ys_pin/_crop.png'; crop.save(cp)
            p=subprocess.run(['swift',OCR,cp],capture_output=True,text=True)
            t=' / '.join(l.strip() for l in p.stdout.split('\n') if l.strip() and not l.startswith('TXT'))
            print(f'   截图{i-22} 标签@({x:.2f},{y:.2f}) 5x裁剪OCR: {t[:100]}')
