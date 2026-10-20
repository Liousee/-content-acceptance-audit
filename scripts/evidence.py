#!/usr/bin/env python3
"""生成「截图证据清单.xlsx」：截图清单 + 作品字段快照。

用法:
  python3 scripts/evidence.py --work-dir ./PO0000000000
  python3 scripts/evidence.py --work-dir ./PO0000000000 --out ./截图证据清单.xlsx

输入:
  <work-dir>/work_rows.json   （由 audit_pipeline.py 生成）
  ~/.cache/yanshou_ocr.json    （可选；用于逐张截图的 OCR 字段）

输出:
  <work-dir>/截图证据清单.xlsx
    · 截图清单：每张截图的 URL/本地路径、类型、OCR 已读字段、数值
    · 作品字段快照：阶段二「内容是否一致」比对的基线
"""
import argparse
import json
import os
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)
from audit_pipeline import parse_ocr  # noqa: E402
from screenshot_coverage import classify  # noqa: E402

HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name='微软雅黑', size=10)
LINK_FONT = Font(name='微软雅黑', size=9, color='0563C1', underline='single')
ALT_FILL = PatternFill('solid', fgColor='F2F7FB')
THIN = Side(style='thin', color='B0B7C3')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def load_json(path, default=None):
    if not path or not os.path.exists(path):
        return default
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def load_ocr_cache(path):
    data = load_json(path, {}) or {}
    now = datetime.now().timestamp()
    ttl = 7 * 86400
    return {k: v for k, v in data.items()
            if isinstance(v, dict) and v.get('ts', 0) > now - ttl}


def row_images(x):
    imgs = x.get('imgs')
    if imgs:
        return list(imgs)
    return [s for s in (x.get('s1'), x.get('s2'), x.get('s3')) if s]


def ocr_of(cache, shot):
    entry = cache.get(shot) or {}
    if 'txt' in entry:
        fin, like, cmt, rep, fav, play, bad = parse_ocr(entry['txt'])
        return {'finish': fin, 'like': like, 'cmt': cmt, 'rep': rep, 'fav': fav,
                'play': play, 'status_bad': sorted(bad), 'txt': entry['txt']}
    return entry


def pct(v):
    return f'{v:.2f}%' if isinstance(v, (int, float)) else None


def style_header(ws, ncol, widths, height=28):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[1].height = height
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(ncol)}{ws.max_row}'


def style_body(ws, r, ncol, center_cols, link_cols=()):
    for c in range(1, ncol + 1):
        cell = ws.cell(r, c)
        cell.border, cell.font = BORDER, BODY_FONT
        cell.alignment = CENTER if c in center_cols else LEFT
        if c in link_cols and isinstance(cell.value, str) and cell.value.startswith(('http://', 'https://')):
            cell.hyperlink, cell.font = cell.value, LINK_FONT
        if r % 2 == 1:
            cell.fill = ALT_FILL


def build(work_dir, out_path, ocr_cache):
    rows_path = os.path.join(work_dir, 'work_rows.json')
    rows = load_json(rows_path)
    if rows is None:
        raise SystemExit(f'缺少 {rows_path}；先运行 audit_pipeline.py 生成')
    cache = load_ocr_cache(ocr_cache)

    wb = Workbook()

    # ---------- 截图清单 ----------
    ws = wb.active
    ws.title = '截图清单'
    hdr = ['序号', '姓名', '昵称', '截图序号', '截图URL/路径', '截图类型', 'OCR已读字段',
           '截图播放量', '赞', '评', '转', '藏', '完播率']
    for c, h in enumerate(hdr, 1):
        ws.cell(1, c, h)
    r = 2
    n_shots = 0
    for x in rows:
        for i, shot in enumerate(row_images(x), 1):
            p = ocr_of(cache, shot)
            got = [k for k in ('play', 'like', 'cmt', 'rep', 'fav', 'finish')
                   if p.get(k) is not None]
            kind = classify(p['txt']) if p.get('txt') else '无OCR文本'
            vals = [x.get('seq'), x.get('name'), x.get('nick'), i, shot, kind,
                    '/'.join(got), p.get('play'), p.get('like'), p.get('cmt'),
                    p.get('rep'), p.get('fav'), pct(p.get('finish'))]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v)
            style_body(ws, r, len(hdr), {1, 4, 6, 8, 9, 10, 11, 12, 13}, {5})
            r += 1
            n_shots += 1
    style_header(ws, len(hdr), [6, 11, 16, 8, 44, 18, 16, 11, 8, 8, 8, 8, 10])

    # ---------- 作品字段快照 ----------
    ws2 = wb.create_sheet('作品字段快照')
    hdr2 = ['序号', '姓名', '昵称', '作品标题', '视频时长(秒)', '系统播放量', '截图阅读量',
            '互动率', '完播率', '激励金额(含税)', '验收费用(未税)', '发布链接', '链接有效性',
            '截图张数', '审核1', '审核2', '审核3', '审核4']
    for c, h in enumerate(hdr2, 1):
        ws2.cell(1, c, h)
    for i, x in enumerate(rows):
        rr = i + 2
        finish = x.get('finish_use')
        if finish is None:
            finish = x.get('finish')
        if finish is None and isinstance(x.get('finish_col'), (int, float)):
            finish = round(x['finish_col'] * 100, 2)
        vals = [x.get('seq'), x.get('name'), x.get('nick'), x.get('title') or '',
                x.get('dur'), x.get('play'), x.get('shot'), x.get('interact'),
                finish, x.get('inc'), x.get('fee'), x.get('pub_link'),
                x.get('link_status') or x.get('link_validity') or '未复核',
                len(row_images(x)), x.get('a1'), x.get('a2'), x.get('a3'), x.get('a4')]
        for c, v in enumerate(vals, 1):
            ws2.cell(rr, c, v)
        style_body(ws2, rr, len(hdr2), {1, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17, 18}, {12})
        for c in (8,):   # 互动率
            ws2.cell(rr, c).number_format = '0.00%'
        for c in (9,):   # 完播率
            ws2.cell(rr, c).number_format = '0.00"%"'
        for c in (6, 7, 10, 11):
            ws2.cell(rr, c).number_format = '#,##0'
        ws2.row_dimensions[rr].height = 20
    style_header(ws2, len(hdr2), [6, 11, 16, 40, 12, 12, 12, 10, 10, 13, 13, 30, 22, 9, 8, 8, 8, 8])

    wb.save(out_path)
    return out_path, n_shots, len(rows)


def main():
    ap = argparse.ArgumentParser(description='Generate a screenshot evidence workbook.')
    ap.add_argument('--work-dir', default=os.getcwd(), help='包含 work_rows.json 的工作目录')
    ap.add_argument('--out', default=None, help='输出路径（默认 <work-dir>/截图证据清单.xlsx）')
    ap.add_argument('--ocr-cache', default=os.path.expanduser('~/.cache/yanshou_ocr.json'))
    args = ap.parse_args()

    out = args.out or os.path.join(args.work_dir, '截图证据清单.xlsx')
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    path, n_shots, n_rows = build(args.work_dir, out, args.ocr_cache)
    print(f'-> {path}（{n_shots} 张截图 + {n_rows} 行作品快照）')


if __name__ == '__main__':
    main()
