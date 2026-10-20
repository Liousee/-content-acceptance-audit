#!/usr/bin/env python3
"""补图归属校验工具: 用户补充截图 → 按内容(发布时间+播放量+互动数据)匹配到正确行
用法: python3 verify_supp_images.py <验收表.xlsx> <补图目录> [--out mapping.json]

背景: 补图文件名常与实际行号不符（整体偏移或完全无关），
必须按内容匹配归位，禁止信文件名。输出映射后：归档到 PO<号>/补图核销/（保留"原文件名N"痕迹）+ 路径写回源表截图列。
"""
import argparse, os, re, json, subprocess
from datetime import datetime
import openpyxl

OCR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ocr.swift')
PAT = re.compile(r'(\d{4}-\d{2}-\d{2})\s*(\d{1,2}:\d{2})')

def ocr_text(img):
    """ocr.swift 输出: 'TXT\t<path>\t<文本>'；文本可能内含换行(整段被拆成多行) → 去掉TXT前缀后按行拼接"""
    p = subprocess.run(['swift', OCR, img], capture_output=True, text=True, timeout=180)
    chunks = []
    for line in p.stdout.split('\n'):
        if line.startswith('TXT\t'):
            parts = line.split('\t', 2)
            chunks.append(parts[2] if len(parts) > 2 else '')
        elif line.startswith('ERR\t'):
            continue
        elif line.strip():
            chunks.append(line)
    return '|'.join(c for c in chunks if c)

def parse_img(txt):
    """提取 发布时间 / 播放量 / 赞 / 评 / 分享 / 收藏 / 完播率"""
    segs = [s.strip() for s in txt.split('|') if s.strip()]
    d = {'pub': None, 'play': None, 'wan': False, 'like': None, 'cmt': None, 'rep': None, 'fav': None, 'finish': None}
    for k, seg in enumerate(segs):
        nxt = segs[k+1] if k+1 < len(segs) else ''
        if d['pub'] is None:
            m = PAT.search(seg)
            if m:
                hh, mm = m.group(2).split(':')
                d['pub'] = f'{m.group(1)} {int(hh):02d}:{mm}'
                continue
        if seg == '播放量':
            m = re.match(r'^(\d+)\s*万$', nxt)
            if m: d['play'] = int(m.group(1))*10000; d['wan'] = True
            elif re.match(r'^\d+$', nxt): d['play'] = int(nxt)
        elif seg == '点赞量' and re.match(r'^\d+$', nxt): d['like'] = int(nxt)
        elif seg == '评论量' and re.match(r'^\d+$', nxt): d['cmt'] = int(nxt)
        elif seg == '分享量' and re.match(r'^\d+$', nxt): d['rep'] = int(nxt)
        elif seg == '收藏量' and re.match(r'^\d+$', nxt): d['fav'] = int(nxt)
        elif ('完播' in seg or '元播' in seg) and '5s' not in seg and len(seg) < 10:
            m = re.search(r'(\d+(?:[.,]\d+)?)\s*%', nxt) or re.search(r'(\d+(?:\.\d+)?)\s*%', seg)
            if m: d['finish'] = float(m.group(1).replace(',', '.'))
    return d

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('imgdir'); ap.add_argument('--out', default=None)
    args = ap.parse_args()
    wb = openpyxl.load_workbook(args.xlsx, data_only=True)['验收链接']
    hdr = [wb.cell(1, c).value for c in range(1, wb.max_column + 1)]
    ci = {h: i+1 for i, h in enumerate(hdr) if h}
    rows = []
    for r in range(2, wb.max_row + 1):
        seq = wb.cell(r, ci['序号']).value
        if seq is None: continue
        rows.append(dict(seq=seq, name=wb.cell(r, ci['姓名']).value, nick=wb.cell(r, ci['昵称']).value,
                         pub=wb.cell(r, ci['发布时间']).value, play=wb.cell(r, ci['用户截图阅读量']).value,
                         like=wb.cell(r, ci['用户截图点赞量']).value, cmt=wb.cell(r, ci['用户截图评论']).value,
                         rep=wb.cell(r, ci['用户截图转发']).value, fav=wb.cell(r, ci['用户截图收藏量']).value))
    imgs = sorted(f for f in os.listdir(args.imgdir) if f.lower().endswith(('.png','.jpg','.jpeg')))
    print(f'补图 {len(imgs)} 张, 表内 {len(rows)} 行\n' + '='*80)
    mapping = {}
    for fn in imgs:
        path = os.path.join(args.imgdir, fn)
        d = parse_img(ocr_text(path))
        hits = []
        for x in rows:
            if d['pub'] is None: break
            pub = x['pub'].strftime('%Y-%m-%d %H:%M') if hasattr(x['pub'], 'strftime') else str(x['pub'])[:16]
            if pub != d['pub']: continue
            score = sum(1 for k in ('play','like','cmt','rep','fav') if d[k] is not None and x[k] is not None and d[k] == x[k])
            if score >= 3: hits.append((score, x))
        hits.sort(key=lambda t: -t[0])
        claim = re.search(r'(\d+)', fn)
        claim = int(claim.group(1)) if claim else None
        if hits:
            seqs = [h[1]['seq'] for h in hits]
            flag = '' if claim in seqs else f'  ⚠文件名标{claim}与实际不符'
            print(f'{fn}\n   → 归属 序号{seqs} {[h[1]["name"] for h in hits]} ({hits[0][0]}/5字段命中){flag}')
            for _, x in hits: mapping[str(x['seq'])] = dict(img=path, name=x['name'], nick=x['nick'],
                                                           matched=dict(pub=d['pub'], play=d['play'], like=d['like'], cmt=d['cmt'], rep=d['rep'], fav=d['fav']),
                                                           finish=d['finish'], file=fn, claim=claim)
        else:
            print(f'{fn}\n   ⚠ 未匹配到任何行 (解析: pub={d["pub"]} play={d["play"]} 赞{d["like"]}/评{d["cmt"]}/转{d["rep"]}/藏{d["fav"]}) — 请人工核对')
    miss = [x['seq'] for x in rows if str(x['seq']) not in mapping]
    print('='*80)
    print(f'匹配成功 {len(mapping)} 行归属 (来自 {len(imgs)} 张图)')
    if args.out:
        json.dump(mapping, open(args.out, 'w'), ensure_ascii=False, indent=1); print('映射已写入', args.out)

if __name__ == '__main__':
    main()
