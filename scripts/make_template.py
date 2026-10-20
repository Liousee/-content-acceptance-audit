#!/usr/bin/env python3
"""生成 Profile A 的空白验收审核底表（37 列 + 审核规则 + 填写说明）。

用法:
  python3 scripts/make_template.py --out 验收审核底表模板.xlsx --rows 50
"""
import argparse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    '序号', '姓名', '联系人（可选）', '车型', '车辆标识（可选）', '昵称', '粉丝数量',
    '主页链接', '任务名称', '子任务ID', '任务上线时间', '发布时间', '是否72HR内发布',
    '发布链接', '系统抓取播放量', '用户截图阅读量', '用户截图点赞量', '用户截图评论',
    '用户截图转发', '用户截图收藏量', '互动率', '完播率', '作品时长',
    '用户上传截图1', '用户上传截图2', '用户上传截图3', '内容等级',
    '激励金额（含税）', '验收费用（未税）', '税率', '含税小计',
    '审核1', '审核2', '审核3', '审核4', '完播率(截图OCR)', '人工确认项',
]

WIDTHS = [
    6, 10, 14, 12, 16, 14, 10, 40, 24, 14, 20, 20, 16,
    40, 14, 14, 14, 14, 14, 14, 10, 10, 10,
    36, 36, 36, 10, 14, 14, 8, 12,
    8, 8, 8, 8, 14, 40,
]

HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name='微软雅黑', size=10)
THIN = Side(style='thin', color='B0B7C3')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

RULES_TEXT = """在此填写本项目的审核规则。建议至少写明：
1) 激励档位：金额 → 播放量/互动率/完播率/时长 门槛
2) 验收费用档位：金额 → 对应档位
3) 粉丝数门槛、发布时间窗口
4) 税率与含税小计勾稽规则
5) 缺失数据、0 元、非标准金额的处理方式
修改后将档位写入 config/rules.local.json，并用 --rules 传给 audit_pipeline.py。
"""

HELP_TEXT = """填写说明：
1. 一行一条作品；序号从 1 开始。
2. 截图列填 URL 或本地绝对路径；支持多张截图交叉核验。
3. 互动率 =（点赞 + 评论 + 转发 + 收藏）/ 用户截图阅读量。
4. 完播率填平台显示的百分数；若源数据口径不明，审核时以截图 OCR 为准。
5. 审核1-4 可留空，由流水线回填；完播率(截图OCR)/人工确认项为输出追加列。
6. 不要在本表填写账号密码、Cookie 或其他凭据。
"""


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'


def main():
    ap = argparse.ArgumentParser(description='Generate a blank Profile A acceptance audit workbook.')
    ap.add_argument('--out', default='验收审核底表模板.xlsx')
    ap.add_argument('--rows', type=int, default=50, help='预留数据行数（默认 50）')
    args = ap.parse_args()

    wb = Workbook()
    ws = wb.active
    ws.title = '验收链接'
    for c, h in enumerate(HEADERS, 1):
        ws.cell(1, c, h)
    style_header(ws, len(HEADERS))
    for r in range(2, args.rows + 2):
        ws.cell(r, 1, r - 1).font = BODY_FONT
        ws.cell(r, 1).alignment = CENTER
        for c in range(2, len(HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if c in (8, 14, 24, 25, 26, 37) else CENTER

    rules = wb.create_sheet('审核规则')
    rules['A1'] = '审核规则'
    rules['A1'].font = Font(name='微软雅黑', size=14, bold=True)
    rules['A3'] = RULES_TEXT
    rules['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    rules.column_dimensions['A'].width = 100
    rules.row_dimensions[3].height = 160

    help_ws = wb.create_sheet('填写说明')
    help_ws['A1'] = '填写说明'
    help_ws['A1'].font = Font(name='微软雅黑', size=14, bold=True)
    help_ws['A3'] = HELP_TEXT
    help_ws['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    help_ws.column_dimensions['A'].width = 100
    help_ws.row_dimensions[3].height = 140

    wb.save(args.out)
    print(f'-> {args.out} ({args.rows} 行 × {len(HEADERS)} 列)')


if __name__ == '__main__':
    main()
